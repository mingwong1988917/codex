#!/usr/bin/env python3
from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook


def main():
    parser = argparse.ArgumentParser(description="Validate a drinkware analysis workbook.")
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--hot-sheet", required=True)
    parser.add_argument("--expected-count", required=True, type=int)
    args = parser.parse_args()

    wb = load_workbook(args.workbook, data_only=False)
    errors = []
    if args.hot_sheet not in wb.sheetnames:
        errors.append(f"missing sheet: {args.hot_sheet}")
    else:
        hot = wb[args.hot_sheet]
        count = sum(hot.cell(row, 2).value is not None for row in range(5, hot.max_row + 1))
        if count != args.expected_count:
            errors.append(f"expected {args.expected_count} hot mothers, got {count}")
        if len(hot._images) != count:
            errors.append(f"expected {count} embedded images, got {len(hot._images)}")

    scatter = next((wb[name] for name in wb.sheetnames if "散点" in name), None)
    if not scatter:
        errors.append("missing scatter sheet")
    else:
        headers = [scatter.cell(4, col).value for col in range(1, scatter.max_column + 1)]
        if "累计热度占比" not in headers:
            errors.append("missing cumulative heat share")

    if errors:
        raise SystemExit("\n".join(errors))
    print("validation passed")


if __name__ == "__main__":
    main()
